from pathlib import Path
import re
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
QA=ROOT.parent.parent/'.sync/academic-doc-qa'
QA.mkdir(parents=True,exist_ok=True)
FONT='/System/Library/Fonts/Supplemental/Times New Roman.ttf'

def diagram(name):
    w,h=1600,580
    im=Image.new('RGB',(w,h),'white');d=ImageDraw.Draw(im);font=ImageFont.truetype(FONT,30)
    def box(x,y,width,label):
        d.rectangle((x,y,x+width,y+78),outline='black',width=2)
        bb=d.textbbox((0,0),label,font=font)
        d.text((x+(width-bb[2]+bb[0])/2,y+23),label,font=font,fill='black')
        return (x+width//2,y,x+width//2,y+78)
    def arrow(points):
        d.line(points,fill='black',width=2)
        x,y=points[-1];px,py=points[-2]
        if y>py: tri=[(x,y),(x-6,y-10),(x+6,y-10)]
        elif y<py:tri=[(x,y),(x-6,y+10),(x+6,y+10)]
        elif x>px:tri=[(x,y),(x-10,y-6),(x-10,y+6)]
        else:tri=[(x,y),(x+10,y-6),(x+10,y+6)]
        d.polygon(tri,fill='black')
    if name=='topology':
        box(80,15,680,'Android gallery host');box(840,15,680,'Web interface / BFF')
        box(80,155,1440,'Authenticated API and services')
        box(80,305,430,'PostgreSQL / jobs');box(585,305,430,'Private image store');box(1090,305,430,'Extraction worker')
        box(400,475,800,'External model provider')
        arrow([(760,54),(840,54)])
        arrow([(1180,93),(1180,120),(800,120),(800,155)])
        arrow([(295,233),(295,305)]);arrow([(800,233),(800,305)])
        arrow([(295,383),(295,420),(1305,420),(1305,383)])
        arrow([(1015,344),(1090,344)])
        arrow([(1305,383),(1305,450),(1000,450),(1000,475)])
        arrow([(80,194),(35,194),(35,514),(400,514)])
        d.text((90,440),'Chat',font=font,fill='black')
    elif name=='entities':
        box(400,15,800,'Owner')
        box(80,155,680,'Account');box(840,155,680,'Upload')
        box(80,305,680,'Transaction');box(840,305,680,'Extraction')
        box(80,475,680,'Transaction evidence link');box(840,475,680,'Review item')
        arrow([(800,93),(800,120),(420,120),(420,155)])
        arrow([(800,120),(1180,120),(1180,155)])
        arrow([(420,233),(420,305)]);arrow([(1180,233),(1180,305)])
        arrow([(420,383),(420,475)]);arrow([(1520,194),(1540,194),(1540,430),(1180,430),(1180,475)])
        arrow([(1520,194),(1560,194),(1560,445),(650,445),(650,475)])
    elif name=='ingestion':
        box(400,15,800,'received');box(400,155,800,'extracting')
        box(65,305,450,'needs_review');box(575,305,450,'imported');box(1085,305,450,'rejected')
        box(65,475,450,'User resolution')
        arrow([(800,93),(800,155)])
        arrow([(800,233),(800,270),(290,270),(290,305)])
        arrow([(800,270),(800,305)]);arrow([(800,270),(1310,270),(1310,305)])
        arrow([(290,383),(290,475)])
        d.text((590,495),'Terminal routing states shown; see text for retries.',font=font,fill='black')
    else:
        box(80,15,1440,'Authenticated approval + idempotency key')
        box(80,125,1440,'Owner advisory lock → watermark row lock')
        box(80,235,1440,'Proposal state + hashes + versions + funding')
        box(80,355,680,'Stale or invalid → reject');box(840,355,680,'Valid → apply + audit')
        box(840,475,680,'Commit financial change')
        arrow([(800,93),(800,125)]);arrow([(800,203),(800,235)])
        arrow([(800,313),(800,333),(420,333),(420,355)])
        arrow([(800,333),(1180,333),(1180,355)]);arrow([(1180,433),(1180,475)])
    p=QA/f'{name}.png';im.save(p);return p

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.61);sec.bottom_margin=Inches(.62);sec.left_margin=Inches(.68);sec.right_margin=Inches(.68)
sec.header_distance=Inches(.20);sec.footer_distance=Inches(.22)
for st in doc.styles:
    if st.type==1:
        st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        for b in st.element.xpath('.//w:pBdr'): b.getparent().remove(b)
        if st.element.rPr is not None:
            for f in st.element.rPr.findall(qn('w:rFonts')):
                for a in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
                    f.attrib.pop(qn('w:'+a),None)
normal=doc.styles['Normal'];normal.font.size=Pt(12)
normal.paragraph_format.line_spacing=1.12;normal.paragraph_format.space_after=Pt(7)
normal.paragraph_format.widow_control=True
for n,size in [('Title',23),('Heading 1',13),('Heading 2',12)]:
    st=doc.styles[n];st.font.size=Pt(size);st.font.bold=True
    st.paragraph_format.space_before=Pt(12 if n!='Title' else 0);st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.keep_with_next=True
for n in ['Header','Footer']:
    doc.styles[n].font.size=Pt(9)

def border(p,edge):
    pp=p._p.get_or_add_pPr(); bs=OxmlElement('w:pBdr');b=OxmlElement('w:'+edge)
    for k,v in [('val','single'),('sz','6'),('space','4'),('color','000000')]:b.set(qn('w:'+k),v)
    bs.append(b);pp.append(bs)
h=sec.header.paragraphs[0];h.text='POCKET PAL PROJECT TEAM\tTECHNICAL SPECIFICATION  |  17 OCTOBER 2026';h.paragraph_format.tab_stops.add_tab_stop(Inches(3.25));border(h,'bottom')
f=sec.footer.paragraphs[0];f.text='Pocket Pal  |  Architecture, Algorithms and Security\t';f.paragraph_format.tab_stops.clear_all();f.paragraph_format.tab_stops.add_tab_stop(Inches(6.85));border(f,'top')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)

def para(t,style=None):
    p=doc.add_paragraph(t,style);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    return p

def table(rows):
    t=doc.add_table(rows=0,cols=len(rows[0]));t.autofit=False
    widths=[1.55,2.65,2.94] if len(rows[0])==3 else [2.05,5.09]
    for col,width in zip(t.columns,widths): col.width=Inches(width)
    pr=t._tbl.tblPr
    b=OxmlElement('w:tblBorders')
    for edge in ['top','bottom','left','right','insideH','insideV']:
        e=OxmlElement('w:'+edge)
        for k,v in [('val','single'),('sz','4'),('color','B0B0B0')]:e.set(qn('w:'+k),v)
        b.append(e)
    pr.append(b)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        trpr=t.rows[-1]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:trpr.append(OxmlElement('w:tblHeader'))
        for ci,(cell,text) in enumerate(zip(cells,row)):
            cell.width=Inches(widths[ci]);p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.paragraph_format.line_spacing=1.04
            r=p.add_run(text);r.font.size=Pt(10);r.bold=ri==0
            cp=cell._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'75');e.set(qn('w:type'),'dxa');mar.append(e)
            cp.append(mar)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEEEEE');cp.append(sh)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

lines=(ROOT/'Pocket_Pal_Technical_Documentation_v2.md').read_text().splitlines();i=0
while i<len(lines):
    line=lines[i]
    if not line:i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            rows.append([c.strip() for c in lines[i].strip('|').split('|')]);i+=1
        table(rows);continue
    if line.startswith('```'):
        p=doc.add_paragraph();p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(8);p.paragraph_format.line_spacing=1
        i+=1;code=[]
        while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
        r=p.add_run('\n'.join(code));r.font.name='Courier New';r.font.size=Pt(10)
    elif line.startswith('# '):
        p=doc.add_paragraph(line[2:],'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1.03
    elif line.startswith('@author '):
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(7);r=p.add_run(line[8:]);r.bold=True;r.font.size=Pt(14)
    elif line.startswith('@subtitle ') or line.startswith('@version '):
        p=doc.add_paragraph(line.split(' ',1)[1]);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(4 if line.startswith('@subtitle') else 20)
        for r in p.runs:r.font.size=Pt(11)
    elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
    elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
    elif line.startswith('@abstract '):
        p=para(line[10:]);p.paragraph_format.line_spacing=1.16
        for r in p.runs:r.italic=True
    elif line.startswith('@keywords '):
        p=para(line[10:]);p.runs[0].bold=True;p.runs[0].font.size=Pt(11)
    elif line.startswith('@figure '):
        name,caption=line[8:].split('|',1);p=doc.add_paragraph();p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(diagram(name)),width=Inches(7.0))
        # Supply meaningful alternative text.
        for prop in p._p.xpath('.//wp:docPr'):prop.set('descr',caption)
        p=doc.add_paragraph(caption);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.runs[0].italic=True;p.runs[0].font.size=Pt(10)
    elif line.startswith('[S') or line.startswith('[R'):
        p=para(line);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    elif line.startswith('@equation '):
        p=doc.add_paragraph(line[10:]);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(5);p.paragraph_format.space_after=Pt(9)
    else:para(line)
    i+=1

doc.core_properties.title='Pocket Pal System Architecture and Technical Specification'
doc.core_properties.author='Pocket Pal Project Team'
doc.core_properties.subject='Implementation baseline and technical design'
doc.core_properties.comments=''
out=ROOT/'Pocket_Pal_Technical_Documentation_v2.docx';doc.save(out)
(QA/'artifact.md').write_text('Reference: HOW_TO_WRITE_AN_ACADEMIC_RESEARCH_PAPER.pdf. Letter single column, Times New Roman body 12, title 23, author 14 bold, italic abstract, keywords and introduction on first page. Own monochrome running metadata replaces journal branding; no watermark or journal attribution. Body content is Pocket Pal implementation documentation. All pages must be rendered and inspected.\n')
print(out)
