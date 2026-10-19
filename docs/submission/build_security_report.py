"""Build the security report only; never rebuild the manually edited technical DOCX."""
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent
D=Document(ROOT/'Pocket_Pal_Technical_Documentation_v2.docx')
for child in list(D._element.body):
 if child.tag!=qn('w:sectPr'):D._element.body.remove(child)
for name in ['Header','Footer']:
 D.styles[name].paragraph_format.tab_stops.clear_all()
for s in D.sections:
 for p in s.header.paragraphs:p.clear()
 h=s.header.paragraphs[0];h.text='POCKET PAL PROJECT TEAM\tSECURITY SELF ASSESSMENT';h.paragraph_format.tab_stops.clear_all();h.paragraph_format.tab_stops.add_tab_stop(s.page_width-s.left_margin-s.right_margin,WD_TAB_ALIGNMENT.RIGHT)
 for p in s.footer.paragraphs:p.clear()
 f=s.footer.paragraphs[0];f.text='Pocket Pal  |  Security Self Assessment\t';f.paragraph_format.tab_stops.clear_all();f.paragraph_format.tab_stops.add_tab_stop(s.page_width-s.left_margin-s.right_margin,WD_TAB_ALIGNMENT.RIGHT)
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
for name in ['Normal','Title','Heading 1','Heading 2','Header','Footer']:
 D.styles[name].font.name='Times New Roman';D.styles[name].font.color.rgb=RGBColor(0,0,0)
D.styles['Normal'].font.size=Pt(12)
D.styles['Normal'].paragraph_format.line_spacing=1.12
D.styles['Normal'].paragraph_format.space_after=Pt(7)
for name,size in [('Title',23),('Heading 1',13),('Heading 2',12)]:
 st=D.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.keep_with_next=True
for name in ['Header','Footer']:D.styles[name].font.size=Pt(9)
def par(text,style=None):
 p=D.add_paragraph(text,style);p.paragraph_format.widow_control=True
 if style is None:p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 return p
def table(rows):
 widths=[2.08,5.06] if len(rows[0])==2 else [1.60,2.70,2.84]
 t=D.add_table(rows=0,cols=len(widths));t.autofit=False
 for c,v in zip(t.columns,widths):c.width=Inches(v)
 borders=OxmlElement('w:tblBorders')
 for edge in ['top','left','bottom','right','insideH','insideV']:
  e=OxmlElement('w:'+edge)
  for k,v in [('val','single'),('sz','4'),('color','B0B0B0')]:e.set(qn('w:'+k),v)
  borders.append(e)
 t._tbl.tblPr.append(borders)
 for index,row in enumerate(rows):
  rr=t.add_row();pr=rr._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
  if not index:pr.append(OxmlElement('w:tblHeader'))
  for c,txt,width in zip(rr.cells,row,widths):
   c.width=Inches(width);p=c.paragraphs[0];p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(4);p.paragraph_format.line_spacing=1.03
   run=p.add_run(txt);run.font.name='Times New Roman';run.font.size=Pt(10);run.bold=not index
   if not index:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEEEEE');c._tc.get_or_add_tcPr().append(sh)
 D.add_paragraph().paragraph_format.space_after=Pt(1)
lines=(ROOT/'Pocket_Pal_Security_Self_Assessment.md').read_text().splitlines();i=0
while i<len(lines):
 l=lines[i];i+=1
 if not l:continue
 if l.startswith('|'):
  rows=[[v.strip() for v in l.strip('|').split('|')]]
  while i<len(lines) and lines[i].startswith('|'):
   rows.append([v.strip() for v in lines[i].strip('|').split('|')]);i+=1
  table(rows);continue
 if l.startswith('# '):
  p=par(l[2:],'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=1.03
 elif l.startswith('### '):par(l[4:],'Heading 2')
 elif l.startswith('## '):par(l[3:],'Heading 1')
 elif l.startswith('@author '):
  p=par(l[8:]);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.runs[0].bold=True;p.runs[0].font.size=Pt(14)
 elif l.startswith('@subtitle ') or l.startswith('@version '):
  p=par(l.split(' ',1)[1]);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.runs[0].font.size=Pt(11)
  if l.startswith('@version'):p.paragraph_format.space_after=Pt(18)
 elif l.startswith('@abstract '):
  p=par(l[10:]);p.runs[0].italic=True
 elif l.startswith('@keywords '):
  p=par(l[10:]);p.runs[0].bold=True;p.runs[0].font.size=Pt(11)
 elif l.startswith('@flow '):
  p=par(l[6:]);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.runs[0].font.size=Pt(10);p.runs[0].bold=True
  sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F2F3F5');p._p.get_or_add_pPr().append(sh)
 else:
  p=par(l)
  if l.startswith('[R'):p.alignment=WD_ALIGN_PARAGRAPH.LEFT
D.core_properties.title='Pocket Pal Security Self Assessment Report';D.core_properties.author='Pocket Pal Project Team';D.core_properties.subject='Permission-tier implementation, evidence and known risks';D.core_properties.comments=''
out=ROOT/'Pocket_Pal_Security_Self_Assessment.docx';D.save(out);print(out)
