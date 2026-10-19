"""Rebuild English POC review documents. Run with the bundled document Python runtime."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT = Path(__file__).resolve().parent

def P(text): return ('p',text)
def H(text): return ('h',text)
def T(headers,rows): return ('t',headers,rows)
def B(text): return ('b',text)

technical = [
('Pocket Pal Technical Documentation',[
P('Topic A Personal Finance Assistant | POC review edition | 17 October 2026'),
P('Pocket Pal turns consented payment images and manually entered records into a personal ledger, budget pockets and a conversational finance assistant. Deterministic backend functions calculate money; the language model explains results and proposes changes that require user approval. The prototype does not connect to banks or move real money.'),
H('Demonstrated outcome and scope'),
P('The owner reports that the Galaxy S22 Ultra automatically discovered and submitted approximately ten dated slip images and displayed results after approximately five minutes. The owner considers this sufficient to proceed with POC documentation. This confirms a reported device workflow, not per-image extraction accuracy, duplicate safety or a measured latency distribution [E4].'),
P('The current APK hosts the integrated Pocket Pal web interface. It requires the local web server, backend and USB forwarding. A signed, remotely hosted standalone release is not delivered. Latency improvement is deferred at the owner’s request; all unverified requirements remain visible.'),
T(['Component','Current implementation'],[
['Android','Kotlin WebView host, WorkManager, MediaStore, bundled ML Kit Latin OCR; package dev.finassist.pocketpal; v0.2.0-gallery reported installed by owner.'],
['Web','Next.js 15 and React 19; Chat, Pockets, Activity, Goals, setup and Settings; same-origin backend proxy.'],
['Backend','Python 3.12 or later, FastAPI, Pydantic, psycopg and PostgreSQL; local evidence includes PostgreSQL 18.4.'],
['Models','Configured Anthropic IDs: claude-haiku-5-5 for extraction and claude-sonnet-5-5 for chat. These are repository configuration values, not a statement about current provider availability.'],
['Storage','PostgreSQL and owner-restricted local upload storage in the current demonstration. Cloud deployment remains proposed.']]),
H('How to read this document'),
P('This is a team review draft for competition technical documentation. Sections cover architecture, algorithms, API and data contracts, reproduction, and evidence. Evidence IDs refer to repository paths in the final section. Source authority is the original product brief, followed by the supplied competition PDF; design documents do not replace their requirements.')]),
('Architecture and trust boundaries',[
P('Data flow: permitted Android photos → local date and candidate filters → authenticated web proxy → FastAPI upload validation → PostgreSQL job queue → extraction provider → deterministic validation and duplicate checks → ledger or review queue → financial tools → chat explanation or pending proposal.'),
H('Client and API boundary'),
P('The web server converts its HttpOnly session cookie into a backend bearer token. The browser never receives the model key. Mutating proxy requests require matching origin, and proxy paths and request sizes are bounded. Android uses an origin-allowlisted, main-frame WebMessage bridge for gallery controls. The old Compose prototype is a separate module and package; its Keystore and queue tests are not automatically evidence for this WebView client.'),
H('Service and domain separation'),
P('backend/src/finassist/domain contains standard-library calculation and validation logic. FastAPI routes establish authenticated ownership; services implement SQL transactions and business operations. Database, storage and provider adapters remain outside the domain. PostgreSQL row-level security and owner-aware constraints add protection beneath application scoping.'),
P('Uploads and their jobs are recorded transactionally. Workers claim jobs using leases and tokens; stale workers cannot persist results after their lease has been replaced. Extraction retries are bounded. Image purge jobs retain retry state instead of permanently abandoning deletion. Upload acceptance therefore precedes extraction, review and import.'),
T(['Boundary','Data and authority'],[
['Phone to local host','Candidate image, capture timestamp and expected owner assertion; authenticated session remains authoritative.'],
['Backend to Anthropic','Consented image for extraction, or bounded conversation and financial tool data for chat. Image text cannot grant permissions.'],
['Model to tools','Validated read, simulation or proposal requests. No approve, delete, consent-change or payment tool.'],
['User to approval service','Explicit approval of an exact proposal, checked against current state and committed in one transaction.']]),
H('Storage responsibilities'),
P('Core records include users and hashed sessions, accounts and balance observations, income schedules, transactions, bills, pockets and budget plans, goals, uploads and review items, jobs, consent grants, proposals, chat messages and audit events. Images are served through owner checks rather than public file paths. Chat retention is optional and distinct from transaction storage.'),
P('Implementation anchors: backend/src/finassist/api/routes.py, services/, infra.py, worker.py and db/migrations/; web/src/app/api/; android/pocketpal/src/main/kotlin/dev/finassist/pocketpal/.')]),
('Image discovery extraction and review',[
H('Discovery and candidate filtering'),
P('The scanner reads accessible MediaStore metadata, distinguishing undated, old, future and in-window images. Eligibility uses DATE_TAKEN within a rolling 30-day interval plus the user’s consent window. DATE_ADDED is not substituted. Full or selected-photo permission determines the visible set. JPEG and PNG inputs are bounded to 8 MiB.'),
P('Filter version 2 uses local Latin OCR. A candidate matches either the original English payment-word rule or a layout rule: decimal money, a masked account or long alphanumeric reference, and an HH:MM time or Buddhist-era year. This is an engineering heuristic, not a trained transaction classifier or a validated privacy boundary. Nonfinancial images can match and true slips can be missed. The eight policy tests use synthetic text, not an image accuracy benchmark [E3].'),
P('Each scan allows at most 100 local checks and ten upload attempts; each owner has a local daily cap of twenty attempts. Attempts are reserved before sending. Versioned rejection markers allow older-filter results to be inspected locally again; accepted metadata and content hashes avoid resending known accepted images. Cache migration still needs broader device regression evidence.'),
H('Extraction and ledger decisions'),
P('The backend validates ownership, consent, size and image type before scheduling extraction. Schema-constrained extraction identifies document type, amount, date, time, counterparty, status and reference where available. Parsing and validation then decide whether the record can be imported or must await review. Processing permission alone does not grant auto-import permission.'),
P('An unresolved amount, direction or payment status does not contribute to financial totals. Review exposes server-defined allowed decisions and candidate records. A receipt can attach to an existing purchase without adding a second expense. A bill remains a bill until a payment is linked. Manual confirmation of an undated image is separate from automatic scanning.'),
H('Known reference limitation'),
P('The active extraction contract has a single untyped bank_ref. Two printed identifiers on one slip may represent different fields. REF-GAP-1 can admit a duplicate when different fields are extracted; REF-GAP-2 can merge different payments when an untyped identifier and amount coincide. Both remain strict expected failures. A typed multi-reference contract is proposed in design document 14, not implemented. General duplicate safety is not claimed [E5].'),
H('Diagnosis on the phone'),
P('Scan monitor separates metadata census, already uploaded/filtered records, consent exclusions, unreadable images, candidate matches and server acceptance. PocketPalScan logs contain fixed stages and aggregate counters rather than OCR or image content. Activity shows processing outcomes; upload counts alone do not establish completed OCR or import.')]),
('Financial algorithms and conversation',[
H('Balance and cash flow'),
P('All monetary records use integer satang. The starting balance is an observation at a specified date or timestamp. Only subsequent eligible transactions adjust its projection; earlier imported history remains available for analysis without subtracting spending already reflected in that balance. A date-only item on the balance date is treated as already reflected under the recorded product policy.'),
P('For the interval from today to, but excluding, the next income date: reserved funds equal eligible unpaid bills plus essentials not already represented by those bills, savings commitments and the buffer. Discretionary funds equal projected available balance minus reserved funds and any simulated purchase. Shortage is the positive magnitude of a negative discretionary balance.'),
P('For positive discretionary funds D and N days, daily allowance is floor(D/N) satang with the remainder recorded; weekly allowance is daily allowance multiplied by min(7,N). Negative discretionary funds display a shortage, never a positive available amount. Percentage buffers identify next income as their base; allocation percentages identify the distributable base.'),
H('Pockets and goals'),
P('Pocket spending is counted within the active plan’s period using category-to-pocket mapping. Categorizing an older transaction changes historical reporting without necessarily changing current-period spending. A pocket transfer redistributes unspent allocations; it neither creates money nor becomes an expense. Budget checks run at proposal creation and again at approval. Bill-linked essentials are reserved once and internal transfer legs remain paired.'),
P('Goal functions use the target, saved amount, deadline, income events and available capacity to calculate required contributions and feasibility. Simulations do not create purchases or alter plans. Proposed plan or goal changes must pass the same approval boundary as chat-generated reallocations.'),
H('Conversation and controlled actions'),
P('The agent can retrieve overview, spending, bills, reviews and insights; simulate a purchase or goal; and propose a budget, goal or reallocation. It uses at most eight tool calls across six model turns. Financial values come from current tools rather than remembered prose. Responses render restricted Markdown and approval cards.'),
P('Tab switching retains owner-bound UI conversation state. With explicit history consent, the server supplies bounded recent context: at most twelve messages and 12000 characters, with a stored-history cap of 200 messages. New chat clears conversation history, not financial records. No embedding database or vector RAG is used. Account changes, stale results and revocation are handled separately from financial data persistence [E1, E2].'),
P('Innovation is the combination of local filtering, uncertainty-aware review, deterministic finance tools and human-approved action. This is an implementation contribution; no research novelty or measured superiority over other assistants is claimed.')]),
('Reproduction and operational workflow',[
H('Local demonstration setup'),
P('Use a Mac with the repository dependencies, Android SDK, Java 17 or later compatible with the project, and an authorized USB-connected S22. Configure the Anthropic secret only in the backend environment. There is no seeded login; create a demonstration account. The gallery profile is separate from dev, e2e, test and realeval.'),
B('cd ~/Development/case/shenzen/backend\nuv run --with pixeltable-pgserver python scripts/dev_server.py --profile gallery --provider anthropic'),
P('In a separate terminal, build the web once after source changes, then start it:'),
B('cd ~/Development/case/shenzen/web\nnpm ci\nnpm run build\nFINASSIST_BACKEND_URL=http://127.0.0.1:8002 npm run start -- --hostname 127.0.0.1 --port 3002'),
P('Build and install the current APK while retaining its data:'),
B('cd ~/Development/case/shenzen/android\nexport ANDROID_HOME="$HOME/Library/Android/sdk"\nexport JAVA_HOME="/Applications/Android Studio.app/Contents/jbr/Contents/Home"\nexport PATH="$ANDROID_HOME/platform-tools:$PATH"\n./gradlew :pocketpal:testDebugUnitTest :pocketpal:assembleDebug\nadb install --user 0 -r pocketpal/build/outputs/apk/debug/pocketpal-debug.apk\nadb reverse tcp:3002 tcp:3002'),
H('Demonstration sequence'),
P('Sign in, set balance and income, grant processing and optional auto-import consent, then open Settings → Slips from the last 30 days. Grant photo access, explicitly enable scanning and select Scan now. Observe the monitor and refresh Activity. Resolve uncertain fields, inspect totals, ask Chat about spending and approve a clearly explained budget proposal only if desired.'),
P('For test copies pushed by adb, MediaStore may require a rescan before the app can discover them. The owner’s device used: adb shell content call --method scan_volume --uri content://media --arg external_primary. This is device test preparation, not a production app requirement or permission bypass. Confirm capture metadata without inventing dates.'),
P('Do not use --reset against gallery or dev, clear app data, or remove original photos. Unit and integration tests must use disposable isolated profiles. Background scheduling depends on Android; a force-stopped app must be opened again. A standalone release requires hosted HTTPS services and deployment verification not provided by this runbook.')]),
('Requirements and evidence register',[
T(['Product feature','Current evidence boundary'],[
['1 Account and setup','Implemented; automated identity/ownership coverage. Final onboarding completeness needs team review.'],
['2 Automatic scanning','Owner reports successful dated-image flow; undated slips excluded, mixed-gallery accuracy and full device acceptance open.'],
['3 Manual transactions','Backend CRUD and validation coverage; complete current web/mobile flow acceptance not established here.'],
['4 Pockets','Allocation/spending and approval-backed changes implemented; complete CRUD and arbitrary editing coverage requires review.'],
['5 Everyday budgeting','Deterministic calculations and proposal checks tested; recommendation quality not generally benchmarked.'],
['6 Goals','Backend feasibility and proposals implemented; complete device interaction acceptance pending.'],
['7 Bills and insights','Backend linking and reporting tested; real gallery-to-insight reconciliation pending.'],
['8 Chat','Connected tools and consented context; scripted tests establish controls, not broad model quality.'],
['9 Security and privacy','Controls implemented with gaps; see companion self-assessment. Deployed encryption not evidenced.'],
['10 Storage and backend','Relational records, jobs and calculation services implemented; reference gaps remain open.'],
['11 Gamification','Optional; not claimed as delivered.']]),
H('Evidence identifiers'),
P('All paths below are relative to evidence/runs/. Counts apply to the recorded source snapshot, not automatically to later edits.'),
P('E1 20261017T012325Z-conversation-gallery-backend/manifest.json: 300 passed; two strict reference xfails serialized as skipped. Real PostgreSQL and storage; fake extraction and scripted chat. E2 20261017T012328Z-conversation-browser/manifest.json: five mocked conversation checks.'),
P('E3 20261017T020449Z-pocketpal-scanpolicy-v2-cloud/: eight synthetic-string policy tests reported by Claude. E4 20261017T040000Z-pocketpal-gallery-dated-scan-owner/summary.md: owner report only, approximate timestamp, no successful-scan logcat captured. E5 20261016T122231Z-reference-contract-regressions/: 57 passed and two open strict xfails.'),
P('E6 20261016T120301Z-real-slip-r2-score/manifest.json: five-image unseen pilot; core fields 5/5 but bank_ref exact match 0/5 with later owner-reported field ambiguity. This is not a general OCR accuracy estimate. E7 20261017T015252Z-native-controls-browser/: pre-v2 mocked monitor checks and APK install record; does not validate the later v2 monitor.'),
H('Submission completion'),
P('Original sources: docs/We go shenzhen case (1).docx and the supplied International Track competition PDF, pages 2–3. The team must still finalize the source revision and GitHub link, presentation, sanitized sandbox execution logs and demonstration video. No competition score is inferred from test counts. This document does not close CP3b.')])]

security=[
('Pocket Pal Security Self Assessment',[
P('Topic A Personal Finance Assistant | POC review edition | 17 October 2026'),
P('This self-assessment describes implemented permission boundaries, repository evidence and known risks for Pocket Pal. It is intended for the team and competition reviewers. It is not an independent penetration test, security certification or legal compliance determination.'),
H('Assessment conclusion'),
P('The prototype separates model proposals from user-authorized commits, authenticates owner-scoped access and applies independent processing and import consent. Automated evidence exercises important negative cases. The owner reports that consented automatic scanning works on the S22 Ultra. These findings support a controlled POC, not unrestricted public release.'),
P('Material open risks include duplicate-reference ambiguity, unmeasured local-filter false positives and negatives, incomplete native lifecycle testing, and unverified deployed encryption and retention. The local debug deployment relies on a trusted Mac and USB connection. Production launch requires remediation and additional evidence.'),
H('Assessed boundary'),
P('The current path is the dev.finassist.pocketpal WebView APK, integrated Next.js web proxy, FastAPI services and worker, PostgreSQL and local image storage, plus external Anthropic processing. Earlier Compose controls are historical evidence only where the exact shared component is still used. In particular, this report does not claim the WebView cookie uses the old Compose Keystore token implementation.'),
T(['Asset','Primary exposure'],[
['Credentials and sessions','Account takeover or session theft.'],
['Slip images and financial records','Sensitive personal and transaction data, including third-party counterparty information.'],
['Budgets and goals','Unauthorized change or misleading financial projections.'],
['Model and database secrets','Service compromise or unbounded provider use.'],
['Audit and execution evidence','Accidental disclosure through logs, screenshots or repository publication.']]),
H('Evidence method'),
P('Claims are grounded in current source inspection and existing recorded runs. No new paid calls, personal-image inspection or penetration tests were performed to author this report. Mocked model tests establish application controls only. Owner observations are labeled separately from automated measurements. The team remains responsible for final risk acceptance and public-release decisions.')]),
('Permission tiers and enforcement',[
P('Permissions are enforced by application services and database transactions, not granted by model text. A tier describes an operation’s authority, not an unrestricted role that the model can assume.'),
T(['Tier','Allowed actor and action','Enforcement'],[
['T0 Read','User or assistant reads the user’s records.','Verified session owner, service scoping and PostgreSQL RLS.'],
['T1 Simulate','Assistant or user calculates affordability and goal alternatives.','Pure calculations; no purchase or plan committed.'],
['T2 Propose','Assistant or user creates a pending budget or goal proposal.','Validated arguments; pending proposal and card only.'],
['T3 Commit','User explicitly approves an exact proposed change.','Owner, hash, version, watermark and expiry checks in a locked transaction.'],
['T4 Consented automation','Worker processes images and imports eligible records.','Separate processing and auto-import grants; recheck before persistence. No budget or goal authority.'],
['T5 Account control','User grants or revokes consent, clears history or requests supported data deletion.','Authenticated user endpoints; absent from the model tool registry.'],
['Prohibited','Real bank payments or transfers.','No bank execution integration or model tool.']]),
H('Approval integrity'),
P('The approval service binds changes to the owner and exact proposal state. Ledger/watermark locks prevent approval from racing with relevant financial changes. Stale proposals are rejected and must be refreshed. Repeated approval cannot apply the same business effect twice. User-controlled edits and approval cards are the authority; a conversational “approved” statement has no committing power.'),
H('Identity and client isolation'),
P('Passwords use Argon2id. Backend sessions are random opaque tokens stored as hashes with a seven-day expiry. owner_id is derived from the verified session. An upload expected_owner_id is an assertion checked against that identity, never a source of authority. Owner-restricted image access and composite constraints supplement row-level security.'),
P('The web cookie is HttpOnly and SameSite Strict; Secure is enabled for HTTPS. Mutations require the expected origin. The native bridge accepts only its allowlisted origin and main frame. Scanner status is hidden when the session binding changes, and logout pauses scanning. Complete device race/revocation testing is still pending.')]),
('Threats controls and evidence',[
T(['Threat','Control and evidence','Residual limitation'],[
['Cross-user data access','Session-derived ownership, RLS, owner-aware constraints and negative API tests [S1].','Coverage is limited to exercised routes and DB roles; no independent penetration test.'],
['Model injection or tool misuse','Fixed extraction schema, validated tool inputs, T0–T2 registry; no model approve/delete/consent tools [S1].','Explanations and proposals can still be misleading; authorization does not prove semantic correctness.'],
['Replay or concurrent changes','Idempotency, constraints, proposal hashes, state versions and transactional locks [S1].','Business duplicate matching remains affected by REF-GAP-1 and REF-GAP-2 [S2].'],
['Revocation and worker crashes','Consent rechecked before persistence; lease token guards; retrying purge jobs [S1].','An already submitted provider request cannot be recalled; device lifecycle coverage incomplete.'],
['Hostile or oversized files','Bounded JPEG/PNG upload handling and schema validation; image processing outside domain [S1].','Image library vulnerabilities and broader malicious-file testing remain operational concerns.'],
['Unsafe chat rendering','Restricted Markdown, raw HTML disabled and controlled links; origin checks at proxy. Source inspection; see implementation sources.','Does not establish complete browser XSS or dependency security.'],
['Sensitive operational logs','Safe exception logging; scanner stages/counters omit OCR, paths and image content [S1, S4].','Audit records still contain business identifiers; access and retention need operational controls.'],
['Excess provider use','Bounded native attempts, bounded agent tool loop and usage-ledger handling [S1].','Local attempt limits are not global billing enforcement or a defense against every compromised client.']]),
H('Authentication and hosting gaps'),
P('No MFA, complete credential-abuse assessment, distributed rate-limit deployment, independent dependency audit or recovery drill is evidenced in this report. These must not be represented as implemented protections simply because earlier design documents proposed them.'),
H('What the automated suite proves'),
P('S1 records real database and application execution with fake extraction and scripted chat. Its 300 passing cases and two strict expected failures are not a security score. Negative tests include ownership, stale/replayed approvals, consent races, paired transfer constraints and safe logging. They do not establish provider behavior, mobile permission lifecycle correctness or production infrastructure security.')]),
('Privacy data lifecycle and known risks',[
H('Consent and external processing'),
P('Android photo access, photo processing, auto-import and chat history retention are separate decisions. Processing consent includes a date window. The local candidate filter reduces uploads but does not guarantee relevance. Anthropic receives selected image content for extraction and financial/conversation data needed for chat. A local result that is later rejected by the backend may already have crossed the external-processing boundary.'),
P('Undated images are excluded from automatic scanning. A manually selected undated image requires explicit confirmation. The dated test copies used on the S22 had metadata recovered by the owner from iPhone Photos; the application did not infer or fabricate capture dates. Whether Android bank apps normally save usable capture metadata remains a product validation question.'),
H('Storage deletion and retention'),
P('Image access is owner-restricted. Deletion removes access and schedules physical purge; a storage failure must remain visible as pending deletion rather than reported success. Chat history is optional, bounded and removable, with revocation protected against in-flight turns repopulating it. Source photos on the phone are outside backend deletion.'),
P('The current local demonstration does not supply evidence of encrypted database volumes, encrypted image storage, backup expiry, provider-side deletion or timed image/audit retention enforcement. Proposed retention durations in design documents are not service guarantees. The team must approve and implement retention periods and disclose external-provider terms before public deployment.'),
T(['Risk','Severity and disposition'],[
['R01 Reference ambiguity','High integrity risk. Alternate or reused untyped identifiers can duplicate or merge records. Open; typed mapping and pipeline regressions required.'],
['R02 Irrelevant image upload','High privacy risk. Layout filter can match nonfinancial content. FP/FN image benchmark and mixed-gallery review pending.'],
['R03 Missing capture dates','Medium functional risk. Valid bank slips may never enter automatic scanning. Preserve current conservative policy; product decision pending.'],
['R04 Device lifecycle races','High potential privacy risk. Revocation, account switching, process death and background scheduling need complete v2 device tests.'],
['R05 Local deployment protections','High release risk. Debug loopback HTTP, unverified encryption at rest and no hosted release assessment. Limit to controlled POC.'],
['R06 Latency and incomplete views','Medium usability risk. Approximate five-minute owner observation; no timing breakdown. Monitor stages; optimization deferred.'],
['R07 Retention and operations','High release risk. Provider retention, backup/deletion policy, abuse controls and incident procedures need confirmation and evidence.']])]),
('Verification references and release decisions',[
H('Evidence register'),
P('S1 evidence/runs/20261017T012325Z-conversation-gallery-backend/manifest.json and junit.xml: 300 passed; two reference strict xfails shown as skipped. Real PostgreSQL 18.4 and local storage; fake extraction and scripted chat. S2 evidence/runs/20261016T122231Z-reference-contract-regressions/: two open reference gaps reproduced; design proposal in docs/design/14-reference-contract.md.'),
P('S3 evidence/runs/20261017T012328Z-conversation-browser/: five mocked conversation checks. S4 evidence/runs/20261017T015252Z-native-controls-browser/: four pre-v2 mocked monitor/bridge checks, with update installation record. S5 evidence/runs/20261017T020449Z-pocketpal-scanpolicy-v2-cloud/: eight synthetic-string policy tests; no image precision/recall measurement.'),
P('S6 evidence/runs/20261017T040000Z-pocketpal-gallery-dated-scan-owner/summary.md: owner-reported approximate ten uploads and five-minute result delay. No successful-run logcat, per-image correctness review or final ledger reconciliation is included. This folder contains an owner summary rather than an automated manifest.'),
H('Before public release'),
P('Close the reference gaps with typed evidence and real pipeline tests; quantify mixed-gallery false positives and negatives; exercise selected/full/revoked permission, logout/account switching, process death and retry on the actual client. Verify TLS, encrypted storage, backups and deletion in the intended deployment. Add and test abuse controls, confirm external processing and retention disclosures, and produce a reproducible sanitized execution package.'),
H('POC acceptance and remaining team decisions'),
P('The owner has accepted the observed device scan as sufficient to start documentation and has deferred latency optimization. This decision does not waive required product features or close CP3b. The team should record risk owners, remediation priority, accepted demonstration constraints and the source revision used for submission.'),
P('The competition requires permission-tier implementation and a known-risk list. This report provides those artifacts with evidence boundaries. It makes no claim of PDPA compliance, OWASP certification, zero vulnerabilities or a predicted competition score.'),
H('Implementation sources'),
P('Primary product source: docs/We go shenzhen case (1).docx. Submission source: supplied International Track competition PDF, page 3. Code anchors: backend/src/finassist/infra.py; services/core.py, agent.py, imports.py and budget.py; worker.py; db/migrations/; web/src/app/api/; android/pocketpal/src/main/kotlin/dev/finassist/pocketpal/. Supporting designs: 03, 06, 14, 15 and 16.'),
P('Review status: prepared by Codex for human team review on 17 October 2026 using repository source and recorded evidence. No independent assessor sign-off is implied. Later implementation changes require this report and its evidence references to be updated.')])]

def build(name,pages):
 d=Document(); sec=d.sections[0];sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
 sec.top_margin=sec.bottom_margin=Inches(.65);sec.left_margin=sec.right_margin=Inches(.7)
 for st in ['Normal','Title','Heading 1','Heading 2','Header','Footer']:
  style=d.styles[st];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
 n=d.styles['Normal'];n.font.size=Pt(10.5);n.paragraph_format.space_after=Pt(7);n.paragraph_format.line_spacing=1.06
 d.styles['Title'].font.size=Pt(25);d.styles['Heading 1'].font.size=Pt(19);d.styles['Heading 2'].font.size=Pt(12)
 header=sec.header.paragraphs[0];header.text='POCKET PAL  |  POC TEAM REVIEW  |  17 OCTOBER 2026';header.style='Header';header.runs[0].font.size=Pt(8)
 foot=sec.footer.paragraphs[0];foot.text='Review draft  •  ';fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld);foot.runs[0].font.size=Pt(8)
 md=[]
 for i,(title,blocks) in enumerate(pages):
  if i:d.add_page_break()
  d.add_paragraph(title,style='Title' if i==0 else 'Heading 1');md.append('# '+title+'\n')
  for b in blocks:
   if b[0] in ('p','h','b'):
    p=d.add_paragraph(b[1],style='Heading 2' if b[0]=='h' else 'Normal')
    if b[0]=='b':
     for run in p.runs:run.font.name='Courier New';run.font.size=Pt(8)
     p.paragraph_format.space_after=Pt(8)
    md.append(('## ' if b[0]=='h' else '')+b[1]+'\n')
   else:
    headers,rows=b[1:];tab=d.add_table(rows=1,cols=len(headers));tab.autofit=False
    widths=([1.65,5.22] if len(headers)==2 else [1.25,2.81,2.81])
    for c,w in zip(tab.columns,widths):c.width=Inches(w)
    for c,t in zip(tab.rows[0].cells,headers):c.text=t
    for row in rows:
     for c,t in zip(tab.add_row().cells,row):c.text=t
    for ri,row in enumerate(tab.rows):
     trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
     if ri==0:trpr.append(OxmlElement('w:tblHeader'))
     for ci,cell in enumerate(row.cells):
      cell.width=Inches(widths[ci]);pr=cell._tc.get_or_add_tcPr()
      sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6ED' if ri==0 else ('F5F7F9' if ri%2==0 else 'FFFFFF'));pr.append(sh)
      borders=OxmlElement('w:tcBorders')
      for side in ['top','left','bottom','right']:
       x=OxmlElement('w:'+side);x.set(qn('w:val'),'single');x.set(qn('w:sz'),'4');x.set(qn('w:color'),'D9D9D9');borders.append(x)
      pr.append(borders)
      mar=OxmlElement('w:tcMar')
      for side in ['top','left','bottom','right']:
       x=OxmlElement('w:'+side);x.set(qn('w:w'),'80');x.set(qn('w:type'),'dxa');mar.append(x)
      pr.append(mar)
      for p in cell.paragraphs:
       p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2);p.paragraph_format.line_spacing=1
       for run in p.runs:run.font.size=Pt(9);run.bold=ri==0
    d.add_paragraph().paragraph_format.space_after=Pt(0)
    md += ['| '+' | '.join(headers)+' |','| '+' | '.join(['---']*len(headers))+' |']+['| '+' | '.join(row)+' |' for row in rows]+['']
 for element in (d.styles.element, d.element):
  for border in list(element.iter(qn('w:pBdr'))): border.getparent().remove(border)
 d.core_properties.author='Pocket Pal project team';d.core_properties.title=pages[0][0];d.core_properties.subject='Evidence based POC review draft'
 d.save(ROOT/(name+'.docx'));(ROOT/(name+'.md')).write_text('\n'.join(md))

if __name__=='__main__':
 build('Pocket_Pal_Technical_Documentation',technical)
 build('Pocket_Pal_Security_Self_Assessment',security)
